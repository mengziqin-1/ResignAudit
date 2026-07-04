import os
import re
import pandas as pd
from datetime import datetime

class ContentAnalyzer:
    def __init__(self):
        self.sensitive_keywords = []
        self.cloud_keywords = []
        self.email_keywords = []
        self.MAX_FILE_SIZE = 50 * 1024 * 1024
        self.load_keywords()
    
    def load_keywords(self):
        keywords_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), 'keywords')
        
        sensitive_file = os.path.join(keywords_dir, 'sensitive_keywords.txt')
        cloud_file = os.path.join(keywords_dir, 'cloud_storage_keywords.txt')
        email_file = os.path.join(keywords_dir, 'email_keywords.txt')
        
        if os.path.exists(sensitive_file):
            with open(sensitive_file, 'r', encoding='utf-8') as f:
                self.sensitive_keywords = [line.strip() for line in f if line.strip()]
        
        if os.path.exists(cloud_file):
            with open(cloud_file, 'r', encoding='utf-8') as f:
                self.cloud_keywords = [line.strip() for line in f if line.strip()]
        
        if os.path.exists(email_file):
            with open(email_file, 'r', encoding='utf-8') as f:
                self.email_keywords = [line.strip() for line in f if line.strip()]
    
    def analyze_files(self, file_scan_results, progress_callback=None):
        results = []
        total_files = len(file_scan_results)
        processed = 0
        
        for file_info in file_scan_results:
            if file_info['file_type'] != 'document':
                continue
            
            if file_info.get('file_size', 0) > self.MAX_FILE_SIZE:
                continue
            
            try:
                text = self._extract_text(file_info['file_path'])
                if text:
                    findings = self._match_keywords(text, file_info)
                    if findings:
                        results.extend(findings)
            except Exception:
                pass
            
            processed += 1
            if progress_callback and processed % 10 == 0:
                progress_callback(f"分析文档: {processed}/{total_files}")
        
        return results
    
    def _extract_text(self, file_path):
        ext = os.path.splitext(file_path)[1].lower()
        
        try:
            if ext == '.txt':
                with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                    return f.read()
            elif ext == '.pdf':
                return self._extract_pdf_text(file_path)
            elif ext in ('.doc', '.docx'):
                return self._extract_doc_text(file_path)
            elif ext in ('.xls', '.xlsx'):
                return self._extract_xls_text(file_path)
            elif ext in ('.ppt', '.pptx'):
                return self._extract_ppt_text(file_path)
            elif ext in ('.html', '.htm'):
                with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                    return f.read()
            elif ext == '.csv':
                with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                    return f.read()
            elif ext == '.json':
                with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                    return f.read()
            elif ext == '.xml':
                with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                    return f.read()
            elif ext == '.md':
                with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                    return f.read()
            elif ext == '.rtf':
                return self._extract_rtf_text(file_path)
        except Exception:
            return ""
        
        return ""
    
    def _extract_pdf_text(self, file_path):
        try:
            import fitz
            doc = fitz.open(file_path)
            text = ""
            for page in doc:
                text += page.get_text()
            return text
        except ImportError:
            return ""
        except Exception:
            return ""
    
    def _extract_doc_text(self, file_path):
        try:
            import docx
            doc = docx.Document(file_path)
            text = ""
            for para in doc.paragraphs:
                text += para.text + "\n"
            return text
        except ImportError:
            return ""
        except Exception:
            return ""
    
    def _extract_xls_text(self, file_path):
        try:
            import pandas as pd
            df = pd.read_excel(file_path)
            return df.to_string()
        except ImportError:
            return ""
        except Exception:
            return ""
    
    def _extract_ppt_text(self, file_path):
        try:
            import pptx
            prs = pptx.Presentation(file_path)
            text = ""
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, 'text'):
                        text += shape.text + "\n"
            return text
        except ImportError:
            return ""
        except Exception:
            return ""
    
    def _extract_rtf_text(self, file_path):
        try:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                rtf_content = f.read()
            text = re.sub(r'\\[a-z]+(\s+[^\s]+)?', '', rtf_content)
            text = re.sub(r'[{}]', '', text)
            return text
        except Exception:
            return ""
    
    def _match_keywords(self, text, file_info):
        findings = []
        
        matched_sensitive = []
        for keyword in self.sensitive_keywords:
            if keyword.lower() in text.lower():
                matched_sensitive.append(keyword)
        
        matched_cloud = []
        for keyword in self.cloud_keywords:
            if keyword.lower() in text.lower():
                matched_cloud.append(keyword)
        
        matched_email = []
        for keyword in self.email_keywords:
            if keyword.lower() in text.lower():
                matched_email.append(keyword)
        
        email_pattern = r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}'
        found_emails = re.findall(email_pattern, text)
        
        phone_pattern = r'1[3-9]\d{9}'
        found_phones = re.findall(phone_pattern, text)
        
        ip_pattern = r'\b(?:\d{1,3}\.){3}\d{1,3}\b'
        found_ips = re.findall(ip_pattern, text)
        
        if matched_sensitive:
            findings.append({
                'type': 'sensitive_keyword',
                'file_path': file_info['file_path'],
                'file_name': file_info['file_name'],
                'modify_time': file_info['modify_time'],
                'matched_keywords': matched_sensitive,
                'severity': 'high',
                'description': f"发现敏感关键词: {', '.join(matched_sensitive[:5])}"
            })
        
        if matched_cloud:
            findings.append({
                'type': 'cloud_storage',
                'file_path': file_info['file_path'],
                'file_name': file_info['file_name'],
                'modify_time': file_info['modify_time'],
                'matched_keywords': matched_cloud,
                'severity': 'medium',
                'description': f"发现云存储相关内容: {', '.join(matched_cloud[:5])}"
            })
        
        if matched_email:
            findings.append({
                'type': 'email',
                'file_path': file_info['file_path'],
                'file_name': file_info['file_name'],
                'modify_time': file_info['modify_time'],
                'matched_keywords': matched_email,
                'severity': 'medium',
                'description': f"发现邮箱相关内容: {', '.join(matched_email[:5])}"
            })
        
        if found_emails:
            findings.append({
                'type': 'email_address',
                'file_path': file_info['file_path'],
                'file_name': file_info['file_name'],
                'modify_time': file_info['modify_time'],
                'matched_keywords': found_emails[:10],
                'severity': 'medium',
                'description': f"发现邮箱地址: {', '.join(found_emails[:5])}"
            })
        
        if found_phones:
            findings.append({
                'type': 'phone_number',
                'file_path': file_info['file_path'],
                'file_name': file_info['file_name'],
                'modify_time': file_info['modify_time'],
                'matched_keywords': found_phones[:10],
                'severity': 'low',
                'description': f"发现手机号码: {', '.join(found_phones[:5])}"
            })
        
        if found_ips:
            findings.append({
                'type': 'ip_address',
                'file_path': file_info['file_path'],
                'file_name': file_info['file_name'],
                'modify_time': file_info['modify_time'],
                'matched_keywords': found_ips[:10],
                'severity': 'low',
                'description': f"发现IP地址: {', '.join(found_ips[:5])}"
            })
        
        return findings
    
    def count_keyword_matches(self, findings, keyword_type=None):
        count = 0
        for finding in findings:
            if keyword_type and finding['type'] != keyword_type:
                continue
            count += len(finding['matched_keywords'])
        return count
    
    def get_summary(self, findings):
        summary = {
            'total_findings': len(findings),
            'by_severity': {'high': 0, 'medium': 0, 'low': 0},
            'by_type': {}
        }
        
        for finding in findings:
            summary['by_severity'][finding['severity']] += 1
            summary['by_type'][finding['type']] = summary['by_type'].get(finding['type'], 0) + 1
        
        return summary